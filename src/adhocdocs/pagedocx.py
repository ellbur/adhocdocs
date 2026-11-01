"""Export a literate page as .docx: headings, paragraphs, lists, emphasis,
block quotations and footnoted citations. Quotations are written as
literal text from the document, with the renderer's marks."""

import io
import re

import docx
from docx.shared import Inches, Pt

from . import literate
from .pages import Page, cells_text, layout_of, markdown, prepare

TOKEN = re.compile(r"ADHOCFILL\d+X")


def _strip_tags(s: str) -> list[tuple[str, bool]]:
    """Citation HTML (only <i> is used) as (text, italic) runs."""
    import html
    out, italic = [], False
    for part in re.split(r"(<i>|</i>)", s):
        if part == "<i>":
            italic = True
        elif part == "</i>":
            italic = False
        elif part:
            out.append((html.unescape(part), italic))
    return out


class Writer:
    def __init__(self, page: Page):
        self.page = page
        self.doc = docx.Document()
        self.source, self.fills = prepare(page)
        self.footnotes = []
        self.written = set()

    def quote_runs(self, para, d, r, block, styles=()):
        for p in literate.pieces(r, layout=layout_of(d, r, self.page), inline=not block, styles=styles):
            if "mark" in p:
                para.add_run(p["mark"])
                continue
            run = para.add_run(p["text"])             # python-docx turns \n into a line break
            st = p.get("style", ())
            run.italic = (p["em"] or "i" in st) or None
            run.bold = "b" in st or None
            run.underline = "u" in st or None
            run.font.strike = "s" in st or None
            run.font.superscript = "sup" in st or None
            run.font.subscript = "sub" in st or None

    def picture(self, para, crops):
        from .quoting import DPI
        import pymupdf
        for i, c in enumerate(crops):
            if i:
                para.add_run().add_break()
            width = pymupdf.Pixmap(c.png).width / DPI
            para.add_run().add_picture(io.BytesIO(c.png), width=Inches(min(6.0, width)))
        if any(c.rendered for c in crops):
            from .pages import RENDERED_NOTE
            para.add_run().add_break()
            note = para.add_run(RENDERED_NOTE)
            note.italic, note.font.size = True, Pt(8)

    def table(self, r):
        width = max(len(row) for row in r.cells)
        t = self.doc.add_table(rows=len(r.cells), cols=width)
        t.style = "Table Grid"
        for i, row in enumerate(r.cells):
            for j, cell in enumerate(row):
                t.cell(i, j).text = cell

    def markdown_table(self, tokens, i):
        """A Markdown table, header row bold. Returns the index of table_close."""
        rows = []
        while tokens[i].type != "table_close":
            t = tokens[i]
            if t.type == "tr_open":
                rows.append([])
            elif t.type == "inline":
                rows[-1].append((tokens[i - 1].type == "th_open", t.children))
            i += 1
        table = self.doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
        table.style = "Table Grid"
        for ri, row in enumerate(rows):
            for ci, (head, children) in enumerate(row):
                para = table.cell(ri, ci).paragraphs[0]
                self.inline(para, children)
                if head:
                    for run in para.runs:
                        run.bold = True
        return i

    def cite_runs(self, para, html_cite):
        for text, italic in _strip_tags(html_cite):
            para.add_run(text).italic = italic or None

    def mono(self, para, text):
        for i, line in enumerate(text.split("\n")):
            if i:
                para.add_run().add_break()
            run = para.add_run(line)
            run.font.name, run.font.size = "Courier New", Pt(8)

    def interpretation(self, para, d, r):
        """Claude's reading: labeled as such, never set like a quotation."""
        from . import interp
        from .pages import _doc_only
        want, err = literate.shown_as(self.page.m, d, r)
        if err:
            raise ValueError(f"reference {d.id}: {err}")
        rec = r.interp
        value = interp.value_of(self.page.m, rec, r.ptr)
        source = "".join(t for t, _ in _strip_tags(self.page.cite(_doc_only(r))))
        if d.kind == "interp":
            para.add_run(interp.as_text(self.page.m, rec, value, want))
            if d.opts.get("cite", "inline") != "none":
                para.add_run(f" (Claude’s reading of {source})").italic = True
            return
        label = para.add_run(f"Interpretation by Claude, not the document’s own content: {rec['title']}. "
                             f"Decoded from {source}; output {rec['output']['sha256'][7:15]}.")
        label.italic, label.font.size = True, Pt(8)
        if want == "image":
            if rec["output"]["type"] != "png":
                raise ValueError(f"reference {d.id}: Word can't hold an SVG picture; have the decoder print PNG, "
                                 "or export html or pdf")
            para.add_run().add_break()
            para.add_run().add_picture(io.BytesIO(interp.output_bytes(self.page.m, rec)), width=Inches(6.0))
        elif want == "table":
            header, rows, _, more = interp.table_rows(rec, value)
            allrows = ([header] if header else []) + rows
            width = max(len(x) for x in allrows) if allrows else 1
            t = self.doc.add_table(rows=len(allrows), cols=width)
            t.style = "Table Grid"
            for i, row in enumerate(allrows):
                for j, cell in enumerate(row):
                    t.cell(i, j).text = str(cell)
                    if header and i == 0:
                        for run in t.cell(i, j).paragraphs[0].runs:
                            run.bold = True
            if more:
                self.doc.add_paragraph(f"[{more} more rows not shown]")
        else:
            para.add_run().add_break()
            self.mono(para, interp.as_text(self.page.m, rec, value, want))

    def fill(self, para, token, style):
        self.written.add(token)
        kind, d, r = self.fills[token]
        if r.status == "broken":
            raise ValueError(f"broken reference {d.id}")
        if d.kind in ("interp", "interp-block"):
            if d.kind == "interp-block":
                para.paragraph_format.left_indent = Inches(0.25)
            self.interpretation(para, d, r)
            return
        citation = self.page.cite(r, d.opts.get("signal"))
        if d.kind == "cite":
            self.cite_runs(para, citation)
            return
        shown = self.page.shown(d, r)
        if shown["error"]:
            raise ValueError(f"reference {d.id}: {shown['error']}")
        styles = (shown.get("styles") or ()) if shown["as"] == "rich" else ()
        if d.kind == "block":
            para.paragraph_format.left_indent = Inches(0.5)
            para.paragraph_format.right_indent = Inches(0.5)
            if shown["as"] == "hex":
                from .raw import hexdump_text
                self.mono(para, hexdump_text(r.data, r.select["bytes"][0]))
            elif shown["as"] == "image":
                self.picture(para, shown["crops"])
            elif shown["as"] == "table":
                self.table(r)
            else:
                self.quote_runs(para, d, r, True, styles)
        elif r.kind == "cells":
            para.add_run("“" + cells_text(r) + "”")
        elif r.kind == "bytes":
            run = para.add_run(r.data.hex(" "))
            run.font.name = "Courier New"
        else:
            para.add_run("“")
            self.quote_runs(para, d, r, False, styles)
            para.add_run("”")
        mode = d.opts.get("cite", "inline")
        if mode == "footnote":
            self.footnotes.append(citation)
            para.add_run(str(len(self.footnotes))).font.superscript = True
        elif mode != "none":
            if d.kind == "block":
                c = self.doc.add_paragraph()
                c.paragraph_format.left_indent = Inches(0.5)
                self.cite_runs(c, citation)
            else:
                para.add_run(" (")
                self.cite_runs(para, citation)
                para.add_run(")")

    def inline(self, para, children, style=None):
        bold = italic = False
        for t in children:
            if t.type == "strong_open":
                bold = True
            elif t.type == "strong_close":
                bold = False
            elif t.type == "em_open":
                italic = True
            elif t.type == "em_close":
                italic = False
            elif t.type in ("softbreak",):
                para.add_run(" ")
            elif t.type == "hardbreak":
                para.add_run().add_break()
            elif t.type in ("text", "code_inline"):
                pos = 0
                for m in TOKEN.finditer(t.content):
                    if m.start() > pos:
                        run = para.add_run(t.content[pos:m.start()])
                        run.bold, run.italic = bold or None, italic or None
                    self.fill(para, m.group(), style)
                    pos = m.end()
                if pos < len(t.content):
                    run = para.add_run(t.content[pos:])
                    run.bold, run.italic = bold or None, italic or None
            elif t.type == "footnote_ref":
                para.add_run(f"[{t.meta['id'] + 1}]").font.superscript = True

    def write(self, out):
        title = self.page.meta.get("title")
        if title:
            self.doc.add_heading(str(title), level=0)
        tokens = markdown().parse(self.source)
        lists = []
        i = 0
        while i < len(tokens):
            t = tokens[i]
            if t.type == "heading_open":
                h = self.doc.add_heading(level=int(t.tag[1]))
                self.inline(h, tokens[i + 1].children)
                i += 3
                continue
            if t.type in ("bullet_list_open", "ordered_list_open"):
                lists.append("List Bullet" if t.type == "bullet_list_open" else "List Number")
            elif t.type in ("bullet_list_close", "ordered_list_close"):
                lists.pop()
            elif t.type == "paragraph_open":
                inl = tokens[i + 1]
                content = inl.content.strip()
                if TOKEN.fullmatch(content) and self.fills[content][1].kind in ("block", "interp-block"):
                    self.fill(self.doc.add_paragraph(), content, None)
                else:
                    p = self.doc.add_paragraph(style=lists[-1] if lists else None)
                    self.inline(p, inl.children)
                i += 3
                continue
            elif t.type == "table_open":
                i = self.markdown_table(tokens, i)
            elif t.type == "footnote_block_open":
                self.doc.add_paragraph("Notes")
            elif t.type == "footnote_open":
                inl = next(x for x in tokens[i:] if x.type == "inline")
                p = self.doc.add_paragraph(f"[{t.meta['id'] + 1}] ")
                self.inline(p, inl.children)
                while tokens[i].type != "footnote_close":
                    i += 1
            i += 1
        missing = [d.id for k, (kind, d, _) in self.fills.items() if kind != "link" and k not in self.written]
        if missing:  # never drop a quotation silently: Markdown this writer doesn't handle
            raise ValueError(f"the Word export can't place reference(s) {', '.join(missing)}; "
                             "move them out of unsupported Markdown, or export html or pdf")
        if self.footnotes:
            self.doc.add_paragraph("Citations")
            for n, c in enumerate(self.footnotes, 1):
                p = self.doc.add_paragraph(f"{n}. ")
                self.cite_runs(p, c)
        self.doc.save(str(out))


def write_docx(page: Page, out):
    Writer(page).write(out)
