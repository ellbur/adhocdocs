import io
import mailbox
import zipfile
from pathlib import Path

import openpyxl

from adhocdocs.extract import extract, kind_for
from adhocdocs.ingest import ingest_bytes
from adhocdocs.matter import Matter
from adhocdocs.search import Options, Searcher

CORPUS = Path(__file__).parent.parent / "corpus"
EML = CORPUS / "muller-z" / "2026-03-09_RE_INV-2026-0347_past-due.eml"
DOCX = CORPUS / "muller-z" / "master-supply-agreement-2025-11-03.docx"


def xlsx_bytes():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Q1"
    ws.append(["Invoice", "Amount", "Note"])
    ws.append(["INV-2026/0347", 48612.5, "never received"])
    ws2 = wb.create_sheet("Q2")
    ws2.append(["nothing", None, "here"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx():
    e = extract("book.xlsx", xlsx_bytes())
    layer = e.layers[0]
    assert e.kind == "spreadsheet" and e.metadata["sheets"] == 2
    assert layer.text == "Invoice\tAmount\tNote\nINV-2026/0347\t48612.5\tnever received\fnothing\t\there"
    cells = {s["ref"]: layer.text[s["start"]:s["end"]] for s in layer.segments if s["kind"] == "cell"}
    assert cells["C2"] == "never received" and cells["B2"] == "48612.5"
    assert [s["name"] for s in layer.segments if s["kind"] == "sheet"] == ["Q1", "Q2"]


def test_cell_locators(tmp_path):
    from adhocdocs import annotations, locators
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes()))
    wb["Q2"].title = "Q2 & later"
    buf = io.BytesIO()
    wb.save(buf)
    m = Matter.create(tmp_path / "m")
    doc = ingest_bytes(m, buf.getvalue(), "book.xlsx", {"custodian": "cust:a", "source": {"collection": "c", "path": "book.xlsx"}})
    m = Matter(m.root)

    r = locators.resolve(m, f"doc:{doc}#sheet=Q1&cells=B1:C2")
    assert r["status"] == "exact" and r["text"] == "Amount\tNote\n48612.5\tnever received"
    text = m.objects.get_text(m.state.layer(r["loc"]["layer"]).text)
    assert text[slice(*r["span"])] == "Amount\tNote\nINV-2026/0347\t48612.5\tnever received"
    assert r["loc"]["quote"] == {"exact": r["text"]}
    assert locators.resolve(m, f"doc:{doc}#cells=a2")["text"] == "INV-2026/0347"       # first sheet by default
    assert locators.resolve(m, f"doc:{doc}#cells=D9:E10")["text"] == "\t\n\t"          # past the data: empty

    # sheet names are percent-encoded in the string form
    loc = {"doc": doc, "select": {"sheet": "Q2 & later", "cells": "A1:C1"}}
    s = locators.format(loc)
    assert s.endswith("#sheet=Q2%20%26%20later&cells=A1:C1") and locators.parse(s)["select"] == loc["select"]
    assert locators.resolve(m, s)["text"] == "nothing\t\there"

    for bad in ("sheet=Nope&cells=A1", "cells=A0", "cells=1A"):
        assert locators.resolve(m, f"doc:{doc}#{bad}")["status"] == "broken"
    changed = {**r["loc"], "quote": {"exact": "something else"}}
    assert locators.resolve(m, changed)["status"] == "broken"

    # tags take cell ranges, with the quote filled in
    annotations.write(m, annotations.CLAUDE, {"op": "define-tag", "tag": "key"})
    [ev] = annotations.write(m, annotations.CLAUDE, {"op": "tag", "tag": "key", "target": f"doc:{doc}#sheet=Q1&cells=B2"})
    assert ev["target"]["select"] == {"sheet": "Q1", "cells": "B2"} and ev["target"]["quote"] == {"exact": "48612.5"}


def test_sniffing_zips():
    assert kind_for(None, DOCX.read_bytes()) == "docx"
    assert kind_for(None, xlsx_bytes()) == "xlsx"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("a.txt", "hello")
    assert kind_for(None, buf.getvalue()) == "zip"


def test_zip_and_mbox_become_families(tmp_path):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mail/one.eml", EML.read_bytes())
        z.writestr("book.xlsx", xlsx_bytes())
        z.writestr("__MACOSX/._junk", b"x")
        z.writestr("empty/", b"")
    box_path = tmp_path / "box.mbox"
    box = mailbox.mbox(box_path)
    box.add(EML.read_bytes())
    box.add(EML.read_bytes().replace(b"nothing to pay yet", b"nothing to see"))
    box.close()
    buf2 = io.BytesIO()
    with zipfile.ZipFile(buf2, "w") as z:          # a mailbox inside an archive
        z.writestr("export/box.mbox", box_path.read_bytes())

    m = Matter.create(tmp_path / "m")
    top = ingest_bytes(m, buf.getvalue(), "evidence.zip", {"custodian": "cust:a", "source": {"collection": "c", "path": "evidence.zip"}})
    nested = ingest_bytes(m, buf2.getvalue(), "export.zip", {"custodian": "cust:a", "source": {"collection": "c", "path": "export.zip"}})
    m = Matter(m.root)
    kids = m.state.children[top]
    assert len(kids) == 2
    box_doc = m.state.children[nested][0]
    assert len(m.state.children[box_doc]) == 2

    s = Searcher(m, Options())
    _, items = s.finish(*s.keyword('"never received"'))
    assert len(items) == 4                      # the email in the zip, the workbook, both mailbox messages
    _, items = s.finish(*s.keyword('"nothing to see"'))
    assert len(items) == 1
    header, _ = s.finish(*s.keyword("anything"))
    assert header["coverage"]["empty"] == 3     # the two zips and the mailbox: extracted, no text of their own


def test_failed_extraction_keeps_the_kind(tmp_path, monkeypatch):
    # e.g. a scan when tesseract is missing: the PDF must stay a PDF so it can be OCRed later
    import adhocdocs.ingest as ingest
    from adhocdocs.extract import unextracted_kind

    def broken(name, data):
        raise FileNotFoundError("tesseract")
    monkeypatch.setattr(ingest, "extract", broken)
    m = Matter.create(tmp_path / "m")
    doc = ingest_bytes(m, b"%PDF-1.4 scanned", "scan.pdf", {"source": {"collection": "c", "path": "scan.pdf"}})
    assert m.objects.get_object(m.state.metadata[doc])["kind"] == "pdf"
    assert unextracted_kind("book.xlsx", b"PK\x03\x04") == "spreadsheet"
    assert unextracted_kind("blob", b"\xff\xfe\x00\x81") == "other"
