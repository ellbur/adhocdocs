import json
from pathlib import Path

import pytest

from adhocdocs import literate, locators, pages
from adhocdocs.ingest import ingest_paths
from adhocdocs.matter import Matter

CORPUS = Path(__file__).parent.parent / "corpus"
MEMO = "memo-collection-bellwether-2026-03-16.pdf"
EMAIL = "2026-03-09_RE_INV-2026-0347_past-due.eml"


@pytest.fixture
def m(tmp_path):
    m = Matter.create(tmp_path / "m")
    ingest_paths(m, [CORPUS / "voss-r" / MEMO, CORPUS / "muller-z" / EMAIL], "x", "c")
    return Matter(m.root)


def doc_named(m, name):
    for doc in m.state.live_documents():
        if m.objects.get_object(m.state.metadata[doc]).get("name") == name:
            return doc


def loc_of(m, name, phrase):
    doc = doc_named(m, name)
    layer = m.state.primary_layer(m.state.documents[doc])
    text = m.objects.get_text(layer.text)
    s = text.index(phrase)
    return locators.char_locator(doc, layer.id, text, s, s + len(phrase)), text, s


def write_page(tmp_path, body, refs):
    p = tmp_path / "page.md"
    p.write_text(body + "\n" + literate.format_refs(refs))
    return p


def test_render_inline_block_and_cite(m, tmp_path):
    loc, _, _ = loc_of(m, EMAIL, "the shipment was never received")
    p = write_page(tmp_path, "---\ntitle: T\n---\nShe wrote that :q{#nr}.\n\n::q{#nr cite=none}\n",
                   [{"id": "nr", **loc}])
    html, results = pages.render(pages.Page(m, p))
    assert results == [{"id": "nr", "status": "exact"}]
    assert "the shipment was never received</span>" in html
    assert "(Mar. 9, 2026)" in html and "Email from Zoë Müller" in html
    assert html.count("<blockquote") == 1


def test_alterations(m, tmp_path):
    phrase = "wrote on March 9 that “the shipment was never received.”"
    loc, text, s = loc_of(m, MEMO, phrase)
    w = s + phrase.index("wrote")
    ref = {"id": "a", **loc, "alter": [
        {"replace": [w, w + 1], "with": "W"},
        {"insert": s + phrase.index("the shipment"), "text": "she said"},
        {"sic": s + phrase.index("never") - 1},
        {"emphasis": [s + phrase.index("never"), s + phrase.index("never") + 5]},
        {"lead": "…"}]}
    r = literate.resolve(m, ref)
    assert r.status == "exact"
    ps = literate.pieces(r, inline=True)
    out = literate.plain(ps)
    assert out.startswith("… [W]rote on March 9 that ‘[she said]the shipment was [sic] never")
    assert "‘" in out and "“" not in out                 # nested quotes become single
    assert any(p.get("em") and p["text"] == "never" for p in ps)
    # reading the text pieces back gives the source, unmarked
    assert "".join(p["text"] for p in literate.pieces(r, layout="lines", inline=False) if "text" in p) == text[s + 1:s + len(phrase)]


def test_join_and_bad_alteration(m):
    _, text, _ = loc_of(m, MEMO, "man-\nager")
    s = text.index("man-\nager")
    loc = locators.char_locator(doc_named(m, MEMO), loc_of(m, MEMO, "man-")[0]["layer"], text, s, s + 9)
    r = literate.resolve(m, {"id": "j", **loc, "alter": [{"join": s + 3}]})
    assert literate.plain(literate.pieces(r)) == "manager"
    r = literate.resolve(m, {"id": "j", **loc, "alter": [{"join": s + 1}]})
    assert r.status == "broken" and "line-end hyphen" in r.error
    r = literate.resolve(m, {"id": "j", **loc, "alter": [{"sic": s + 50}]})
    assert r.status == "broken"


def test_elision_and_transcript_free_parts(m, tmp_path):
    a, text, _ = loc_of(m, MEMO, "Bellwether Marine Systems, Inc. has not paid")
    b, _, _ = loc_of(m, MEMO, "which was due on March 7, 2026")
    p = tmp_path / "page.md"
    p.write_text("::q{#e}\n")
    pages.add_ref(m, p, "e", [json.dumps(a), json.dumps(b)])
    with pytest.raises(pages.PageError, match="order"):
        pages.add_ref(m, p, "e2", [json.dumps(b), json.dumps(a)])
    page = pages.Page(m, p)
    assert page.check() == [{"id": "e", "status": "exact"}]
    assert " … " in pages.render(page)[0]


def test_reanchor_and_update(m, tmp_path):
    loc, _, _ = loc_of(m, EMAIL, "the shipment was never received")
    moved = {**loc, "layer": "sha256:" + "0" * 64, "select": {"char": [0, 31]}}
    p = write_page(tmp_path, "X :q{#nr}", [{"id": "nr", **moved}])
    page = pages.Page(m, p)
    assert page.check()[0]["status"] == "re-anchored"
    with pytest.raises(pages.PageError, match="re-anchored"):
        pages.export(page, "html", tmp_path / "out.html")
    assert [c["id"] for c in pages.update_refs(m, p)] == ["nr"]
    assert pages.Page(m, p).check() == [{"id": "nr", "status": "exact"}]


def test_broken_refuses_export_and_never_renders_text(m, tmp_path):
    loc, _, _ = loc_of(m, EMAIL, "the shipment was never received")
    bad = {**loc, "quote": {**loc["quote"], "exact": "the shipment was delivered on time"}}
    p = write_page(tmp_path, "X :q{#nr}", [{"id": "nr", **bad}])
    html, results = pages.render(pages.Page(m, p))
    assert results[0]["status"] == "broken"
    assert "broken reference: nr" in html and "never received" not in html and "delivered on time" not in html
    with pytest.raises(pages.PageError, match="broken"):
        pages.export(pages.Page(m, p), "docx", tmp_path / "out.docx")


def test_lints(m, tmp_path):
    loc, _, _ = loc_of(m, EMAIL, "the shipment was never received")
    p = write_page(tmp_path, '> typed\n\nShe said "we will pay it next week for sure" and :q{#missing}.',
                   [{"id": "nr", **loc}])
    lints = {r.get("lint") for r in pages.Page(m, p).check()}
    assert {"blockquote", "unbacked-quote", "missing-ref", "unused-ref"} <= lints


def test_export_provenance_and_cite_as(m, tmp_path):
    loc, _, _ = loc_of(m, EMAIL, "the shipment was never received")
    (m.workspace / "cite-as.jsonl").write_text(json.dumps({"doc": loc["doc"], "cite": "Müller Email"}) + "\n")
    p = write_page(tmp_path, ":q{#nr cite=footnote}", [{"id": "nr", **loc}])
    for fmt in ("html", "docx", "pdf"):
        res = pages.export(pages.Page(m, p), fmt, tmp_path / f"out.{fmt}")
        prov = json.loads(Path(res["provenance"]).read_text())
        assert prov["quotations"][0]["rendered"] == "the shipment was never received"
        assert prov["quotations"][0]["citation"] == "Müller Email (Mar. 9, 2026)"
    assert sum(1 for e in Matter(m.root).logs["workspace"].read() if e["op"] == "export") == 3


def test_custom_cite_style(m, tmp_path):
    loc, _, _ = loc_of(m, EMAIL, "the shipment was never received")
    (tmp_path / "style.py").write_text("def cite(info):\n    return 'CUSTOM ' + info['kind']\n")
    p = write_page(tmp_path, "---\ncite-style: style.py\n---\n:q{#nr}", [{"id": "nr", **loc}])
    assert "CUSTOM email" in pages.render(pages.Page(m, p))[0]


def test_docx_export_keeps_tables(m, tmp_path):
    import docx
    loc, _, _ = loc_of(m, EMAIL, "the shipment was never received")
    p = write_page(tmp_path, "| Date | Evidence |\n|---|---|\n| 9 Mar | :q{#nr} |", [{"id": "nr", **loc}])
    pages.export(pages.Page(m, p), "docx", tmp_path / "out.docx")
    table = docx.Document(str(tmp_path / "out.docx")).tables[0]
    assert table.cell(0, 0).text == "Date" and table.cell(0, 0).paragraphs[0].runs[0].bold
    assert "the shipment was never received" in table.cell(1, 1).text


def test_docx_export_refuses_to_drop_a_quotation(m, tmp_path):
    loc, _, _ = loc_of(m, EMAIL, "the shipment was never received")
    p = write_page(tmp_path, "See [the email](ref:nr).\n\n```\n:q{#nr}\n```", [{"id": "nr", **loc}])
    with pytest.raises(ValueError, match="can't place reference"):
        pages.export(pages.Page(m, p), "docx", tmp_path / "out.docx")


def test_ref_from_words(m, tmp_path):
    loc, _, _ = loc_of(m, EMAIL, "the shipment was never received")
    # whitespace runs and quote marks may differ from the document; the quote is the layer's
    found = pages.find_words(m, "past-due", ["the  shipment\nwas never received"])
    assert found[0]["select"] == loc["select"] and found[0]["quote"]["exact"] == "the shipment was never received"
    with pytest.raises(pages.PageError, match="occur in 2 documents; pick one with --in"):
        pages.find_words(m, None, ["the shipment was never received"])
    p = write_page(tmp_path, "X :q{#nr}", [])
    pages.add_ref(m, p, "nr", pages.find_words(m, loc["doc"][7:19], ["the shipment", "received"]))
    assert pages.Page(m, p).check()[0]["status"] == "exact"
    assert " … " in pages.render(pages.Page(m, p))[0]
    with pytest.raises(pages.PageError, match="weren't found"):
        pages.find_words(m, "past-due", ["the shipment was delivered on time"])
    with pytest.raises(pages.PageError, match="if case is ignored"):
        pages.find_words(m, "past-due", ["THE SHIPMENT was never received"])
    with pytest.raises(pages.PageError, match="times in"):
        pages.find_words(m, "past-due", ["the"])
