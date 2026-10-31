"""Exhibits: stamping, numbering, the exhibit list, exh: locators and citations."""

import csv
from pathlib import Path

import docx
import pymupdf
import pytest

from adhocdocs import annotations, exhibits, lists, literate, pages
from adhocdocs.ingest import ingest_paths
from adhocdocs.matter import Matter
from adhocdocs.produce import ProductionError, produce

CORPUS = Path(__file__).parent.parent / "corpus"
FILES = {"letter": CORPUS / "voss-r" / "demand-letter-2026-03-25.pdf",
         "invoice": CORPUS / "muller-z" / "INV-2026-0347.pdf",
         "msa": CORPUS / "muller-z" / "master-supply-agreement-2025-11-03.docx"}


@pytest.fixture(scope="module")
def m(tmp_path_factory):
    m = Matter.create(tmp_path_factory.mktemp("exhibits") / "m")
    ingest_paths(m, list(FILES.values()), "x", "c")
    return Matter(m.root)


def doc_of(m, key):
    return next(d for d in m.state.live_documents() if m.objects.get_object(m.state.metadata[d]).get("name") == FILES[key].name)


def list_of(path, items):
    lists.write(path, lists.new("x", [])[0], items)
    return path


def test_exhibits(m, tmp_path):
    letter, invoice, msa = doc_of(m, "letter"), doc_of(m, "invoice"), doc_of(m, "msa")
    # the invoice is produced first, so its exhibit is the Bates-stamped copy
    produce(m, list_of(tmp_path / "p.list.jsonl", [{"loc": "doc:" + invoice}]), "vol001", "HFS")
    first = exhibits.exhibits(m, list_of(tmp_path / "a.list.jsonl", [
        {"loc": "doc:" + letter},
        {"loc": f"doc:{letter}#page=2", "fields": {"description": "Signature page"}},
        {"loc": "doc:" + invoice},
        {"loc": "doc:" + msa}]), "dep-a", witness="Kettering", when="2026-04-28")
    assert (first["first"], first["last"]) == ("Exhibit 1", "Exhibit 4")
    out = Path(first["path"])
    page_one = pymupdf.open(out / "PDF" / "Exhibit_1.pdf")[0].get_text()
    assert "EXHIBIT 1" in page_one and "Kettering" in page_one and "Exhibit 1 · page 1 of 2" in page_one
    assert pymupdf.open(out / "PDF" / "Exhibit_2.pdf").page_count == 1
    assert "HFS000001" in pymupdf.open(out / "PDF" / "Exhibit_3.pdf")[0].get_text()      # the produced copy

    rows = list(csv.DictReader(open(out / "exhibit-list.csv")))
    assert [r["Exhibit"] for r in rows] == ["Exhibit 1", "Exhibit 2", "Exhibit 3", "Exhibit 4"]
    assert rows[1]["Description"] == "Signature page" and rows[2]["Bates"] == "HFS000001" and rows[0]["Date"] == "2026-03-25"
    assert docx.Document(str(out / "exhibit-list.docx")).tables[0].cell(3, 0).text == "Exhibit 3"
    _, items = lists.read(out / "exhibits.list.jsonl")
    assert items[1]["fields"]["exhibit"] == "Exhibit 2" and items[1]["loc"].endswith("#page=2")

    # numbers run on, and are never reused
    second = exhibits.exhibits(m, list_of(tmp_path / "b.list.jsonl", [{"loc": "doc:" + letter}]), "dep-b")
    assert second["first"] == "Exhibit 5"
    with pytest.raises(ProductionError, match="already used"):
        exhibits.exhibits(m, list_of(tmp_path / "c.list.jsonl", [{"loc": "doc:" + letter}]), "dep-c", start=2)
    assert not (m.workspace / "exhibits" / "dep-c").exists()
    dx = exhibits.exhibits(m, list_of(tmp_path / "d.list.jsonl", [{"loc": "doc:" + msa}]), "dep-d", prefix="DX-")
    assert dx["first"] == "DX-1"

    # exh: locators
    hit = exhibits.resolve_exhibit(m, "exh:2")
    assert len(hit) == 1 and hit[0]["exhibit"] == "Exhibit 2"
    with pytest.raises(Exception, match="could be"):
        exhibits.resolve_exhibit(m, "exh:1")                   # Exhibit 1 and DX-1
    assert exhibits.resolve_exhibit(m, "exh:Exhibit1#page=2")[0]["loc"] == {"doc": letter, "select": {"page": 2}}
    assert exhibits.resolve_exhibit(m, "exh:Ex. 2")[0]["loc"] == {"doc": letter, "select": {"page": 2}}
    assert exhibits.resolve_exhibit(m, "exh:dx-1")[0]["loc"]["doc"] == msa
    with pytest.raises(Exception, match="has 1 pages"):
        exhibits.resolve_exhibit(m, "exh:2#page=2")

    # citations: "Ex. N at p" when the page asks for exhibits
    text = m.objects.get_text(m.state.primary_layer(m.state.documents[letter]).text)
    s = text.index("Very truly yours")
    p = tmp_path / "page.md"
    p.write_text("---\ncite-exhibits: dep-a\n---\n:q{#sig}\n\n" + literate.format_refs([
        {"id": "sig", "doc": letter, "layer": m.state.primary_layer(m.state.documents[letter]).id,
         "select": {"char": [s, s + 16]}, "quote": {"exact": text[s:s + 16]}}]))
    html_doc, _ = pages.render(pages.Page(m, p))
    assert "<cite data-ref=\"sig\">Ex. 1 at 1</cite>" in html_doc      # the letter's first page is Exhibit 1's first
    pages.add_ref(m, p, "whole", ["exh:4"])
    assert literate.split(p.read_text())[1][-1]["doc"] == msa


def test_withheld_documents_are_not_exhibits(m, tmp_path):
    msa = doc_of(m, "msa")
    annotations.write(m, annotations.CLAUDE, {"op": "define-tag", "tag": "privileged"})
    annotations.write(m, "user:owen", {"op": "tag", "tag": "privileged", "target": "doc:" + msa})
    produce(m, list_of(tmp_path / "p.list.jsonl", [{"loc": "doc:" + msa}]), "vol-w", "WH")
    with pytest.raises(ProductionError, match="withheld"):
        exhibits.exhibits(m, list_of(tmp_path / "w.list.jsonl", [{"loc": "doc:" + msa}]), "dep-w")
    ok = exhibits.exhibits(m, list_of(tmp_path / "w.list.jsonl", [{"loc": "doc:" + msa}]), "dep-w", originals=True)
    assert ok["exhibits"] == 1
