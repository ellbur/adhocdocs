"""Quoting more than words in pages: formatting (as=rich), pictures of pages
(as=image), and spreadsheet cells (as=table)."""

import io
import json
from pathlib import Path

import docx
import pytest

from adhocdocs import literate, locators, pages
from adhocdocs.extract import TextBuilder, extract
from adhocdocs.extract.simple import html_to_text
from adhocdocs.ingest import ingest_paths
from adhocdocs.matter import Matter

CORPUS = Path(__file__).parent.parent / "corpus"
FILES = {"memo": CORPUS / "voss-r" / "memo-collection-bellwether-2026-03-16.pdf",
         "letter": CORPUS / "muller-z" / "response-to-demand-2026-04-02.docx",
         "mail": CORPUS / "okafor-d" / "2026-02-19_shipped-PO-BMS-26-0112.eml",
         "aging": CORPUS / "voss-r" / "ar-aging-2026-03-31.csv"}


@pytest.fixture(scope="module")
def m(tmp_path_factory):
    m = Matter.create(tmp_path_factory.mktemp("quoting") / "m")
    ingest_paths(m, list(FILES.values()), "x", "c")
    return Matter(m.root)


def doc_of(m, key):
    name = FILES[key].name
    return next(d for d in m.state.live_documents() if m.objects.get_object(m.state.metadata[d]).get("name") == name)


def text_ref(m, key, phrase, rid, extra=0):
    doc = doc_of(m, key)
    layer = m.state.primary_layer(m.state.documents[doc])
    text = m.objects.get_text(layer.text)
    s = text.index(phrase)
    return {"id": rid, **locators.char_locator(doc, layer.id, text, s, s + len(phrase) + extra)}


def test_html_styles_follow_the_text():
    b = TextBuilder()
    b.add("Header\n")
    html_to_text('<p>Plain <b>bold <i>both</i></b> <span style="font-weight:700">css</span> x<sup>2</sup></p>', b)
    got = {b.text[a:e]: set(st) for a, e, st in b.styles}
    assert got == {"bold": {"b"}, "both": {"b", "i"}, "css": {"b"}, "2": {"sup"}}


def test_docx_run_styles():
    d = docx.Document()
    p = d.add_paragraph("Plain ")
    p.add_run("bold").bold = True
    p.add_run(" and ")
    p.add_run("italic").italic = True
    buf = io.BytesIO()
    d.save(buf)
    layer = extract("x.docx", buf.getvalue()).layers[0]
    assert {layer.text[a:e]: set(st) for a, e, st in layer.styles} == {"bold": {"b"}, "italic": {"i"}}


def test_rich_quotes(m, tmp_path):
    refs = [text_ref(m, "letter", "BELLWETHER MARINE SYSTEMS", "letter", 20),
            text_ref(m, "mail", "Bill of lading", "mail", 12),
            text_ref(m, "memo", "MEMORANDUM", "memo", 10)]
    p = tmp_path / "rich.md"
    p.write_text(":q{#letter as=rich} :q{#mail as=rich} :q{#memo as=rich} :q{#memo}\n\n" + literate.format_refs(refs))
    page = pages.Page(m, p)
    html_doc, results = pages.render(page)
    assert all(r.get("status") == "exact" for r in results if "id" in r)
    assert '<b><span data-text=' in html_doc and html_doc.count("adhoc-rich") == 3
    # the words are still the layer's: formatting wraps the same data-text spans
    plain_page = pages.Page(m, p)
    for rid in ("letter", "mail", "memo"):
        r = plain_page.resolved(rid)
        assert literate.plain(literate.pieces(r, styles=plain_page.styles(r)[0])) == literate.plain(literate.pieces(r))
    out = tmp_path / "rich.docx"
    pages.export(page, "docx", out)
    runs = [run for para in docx.Document(str(out)).paragraphs for run in para.runs]
    assert any(run.bold and "BELLWETHER" in run.text for run in runs)
    assert any(run.bold and run.text.startswith("Bill of lading") for run in runs)
    prov = json.loads(out.with_name(out.name + ".provenance.json").read_text())["quotations"]
    assert [q.get("formatting") for q in prov] == ["Word run formatting", "the email's HTML", "PDF fonts (approximate)", None]


def test_pictures_of_pages(m, tmp_path):
    memo = doc_of(m, "memo")
    refs = [text_ref(m, "memo", "MEMORANDUM", "head", 5),
            {"id": "region", "doc": memo, "select": {"page": 1, "xywh": [0.1, 0.05, 0.8, 0.2]}}]
    p = tmp_path / "img.md"
    p.write_text("::q{#head as=image}\n\n::q{#region}\n\n::q{#head as=image highlight=yes}\n\n" + literate.format_refs(refs))
    page = pages.Page(m, p)
    html_doc, results = pages.render(page)
    assert html_doc.count('<img src="data:image/png;base64,') == 3
    html_doc, _ = pages.render(page, standalone=False, images="files")
    assert len(page.images) == 3 and all(png.startswith(b"\x89PNG") for png in page.images.values())
    head = page.shown(literate.directives(page.prose)[0], page.resolved("head"))["crops"][0]
    assert head.page == 1 and "MEMORANDUM" in head.alt and 0 <= head.rect[0] < head.rect[2] <= 1
    out = tmp_path / "img.docx"
    pages.export(page, "docx", out)
    assert len(docx.Document(str(out)).inline_shapes) == 3
    prov = json.loads(out.with_name(out.name + ".provenance.json").read_text())["quotations"]
    assert [q["as"] for q in prov] == ["image", "image", "image"] and prov[1]["locators"][0]["select"]["xywh"]
    assert all(q["images"][0]["sha256"].startswith("sha256:") for q in prov)
    pages.export(page, "pdf", tmp_path / "img.pdf")


def test_cells(m, tmp_path):
    aging = doc_of(m, "aging")
    p = tmp_path / "cells.md"
    p.write_text("x")
    pages.add_ref(m, p, "table", [f"doc:{aging}#cells=A1:C3"])
    pages.add_ref(m, p, "row", [f"doc:{aging}#cells=A2:B2"])
    ref = literate.split(p.read_text())[1][0]
    assert ref["quote"]["exact"].startswith("customer\tinvoice\tinvoice_date\n") and ref["layer"]
    p.write_text("::q{#table}\n\nInline: :q{#row}.\n\n" + p.read_text().split("\n", 1)[1])
    page = pages.Page(m, p)
    html_doc, _ = pages.render(page)
    assert "<table>" in html_doc and "<td>invoice_date</td>" in html_doc
    assert "Bellwether Marine Systems; INV-2026/0347" in html_doc and "cells A2:B2" in html_doc
    out = tmp_path / "cells.docx"
    pages.export(page, "docx", out)
    t = docx.Document(str(out)).tables[0]
    assert t.cell(0, 1).text == "invoice" and len(t.rows) == 3


def test_what_each_kind_allows(m, tmp_path):
    memo, mail = doc_of(m, "memo"), doc_of(m, "mail")
    refs = [text_ref(m, "mail", "Bill of lading", "mail"),
            {"id": "region", "doc": memo, "select": {"page": 1, "xywh": [0.1, 0.1, 0.5, 0.5]}},
            {"id": "bad", "doc": memo, "select": {"page": 1, "xywh": [0.8, 0.1, 0.5, 0.5]}},
            text_ref(m, "memo", "MEMORANDUM", "memo")]
    p = tmp_path / "bad.md"
    p.write_text("::q{#mail as=image}\n\n:q{#memo as=image}\n\n:q{#region as=text}\n\n::q{#bad}\n\n::q{#memo as=table}\n\n"
                 + literate.format_refs(refs))
    errors = {r["id"]: r.get("error", "") for r in pages.Page(m, p).check() if "id" in r}
    assert "paginated" in errors["mail"] and "block" in errors["memo"]
    assert "image" in errors["region"] and "xywh" in errors["bad"]
    with pytest.raises(pages.PageError):
        pages.export(pages.Page(m, p), "html", tmp_path / "bad.html")
