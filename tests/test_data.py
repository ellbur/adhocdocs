"""Files of unknown format or schema: byte locators, JSON pointers, schemas
and interpretations (DATA.md)."""

import json
import struct
import sys

import pytest

from adhocdocs import interp, literate, locators, pages, raw, structured
from adhocdocs.extract import kind_for
from adhocdocs.ingest import ingest_paths
from adhocdocs.matter import Matter

RECORDS = [(1773000000, "Harrowgate", 1250.0), (1773086400, "Bellwether", -310.5), (1773172800, "Zoë Müller", 42.0)]

DECODER = '''import json, struct, sys
b = sys.stdin.buffer.read()
n, = struct.unpack_from("<H", b, 6)
i, out = 8, []
for _ in range(n):
    s = i
    t, l = struct.unpack_from("<IH", b, i); i += 6
    name = b[i:i + l].decode(); i += l
    a, = struct.unpack_from("<d", b, i); i += 8
    out.append({"time": t, "party": name, "amount": a, "_bytes": [s, i]})
print(json.dumps({"records": out}))
'''


def ledger() -> bytes:
    b = b"TXLG\x01\x00" + struct.pack("<H", len(RECORDS))
    for t, n, a in RECORDS:
        nb = n.encode()
        b += struct.pack("<IH", t, len(nb)) + nb + struct.pack("<d", a)
    return b


EXPORT = json.dumps({"account": "PO-BMS-26-0112",
                     "items": [{"sku": "PN 44-B/7", "qty": 400, "note": "Café pallet"}, {"sku": "X", "qty": 2}]},
                    indent=2, ensure_ascii=True)


@pytest.fixture
def m(tmp_path):
    files = tmp_path / "files"
    files.mkdir()
    (files / "ledger.txl").write_bytes(ledger())
    (files / "export.json").write_text(EXPORT)
    (files / "events.jsonl").write_text('{"t":"2026-03-09","msg":"never received"}\n\n{"t":"2026-03-10","msg":"signed"}\n')
    (files / "feed.xml").write_text('<?xml version="1.0"?><orders><order id="1"><sku>A</sku></order>'
                                    '<order id="2"><sku>B</sku></order></orders>')
    (files / "dupkeys.json").write_text('{"a": 1, "a": 2}')
    m = Matter.create(tmp_path / "m")
    ingest_paths(m, [files], None, "c")
    m = Matter(m.root)
    (m.workspace / "decode.py").write_text(DECODER)
    return m


def doc_named(m, name):
    for doc in m.state.live_documents():
        if m.objects.get_object(m.state.metadata[doc]).get("name") == name:
            return doc


def meta(m, name):
    return m.objects.get_object(m.state.metadata[doc_named(m, name)])


# -- recognizing ------------------------------------------------------------------------------

def test_kinds(m):
    assert meta(m, "ledger.txl")["kind"] == "other"
    assert meta(m, "export.json").items() >= {"kind": "json", "format": "json", "valid": True, "shape": "object"}.items()
    assert meta(m, "events.jsonl")["records"] == 2
    assert meta(m, "feed.xml")["kind"] == "xml"
    # by content, without a telling name
    assert kind_for("export", b'{"a": [1, 2]}') == "json"
    assert kind_for("log", b'{"a": 1}\n{"a": 2}\n') == "ndjson"
    assert kind_for("x", b"<?xml version='1.0'?><a/>") == "xml"
    assert kind_for("notes", b"{not json at all") == "text"
    assert kind_for("db", b"SQLite format 3\x00" + bytes(100)) == "other"


def test_binary_metadata_records_signatures(tmp_path):
    (tmp_path / "f").mkdir()
    (tmp_path / "f" / "x.db").write_bytes(b"SQLite format 3\x00" + bytes(200))
    m = Matter.create(tmp_path / "m")
    ingest_paths(m, [tmp_path / "f"], None, "c")
    m = Matter(m.root)
    assert meta(m, "x.db")["signature"] == ["SQLite database"]


def test_signature_search_field(tmp_path):
    import json as _json
    from adhocdocs.index import open_index
    from adhocdocs.search import Options, Searcher
    (tmp_path / "f").mkdir()
    (tmp_path / "f" / "x.db").write_bytes(b"SQLite format 3\x00" + bytes(200))
    (tmp_path / "f" / "y.gz").write_bytes(b"\x1f\x8b" + bytes(200))
    (tmp_path / "f" / "z.txt").write_text("sqlite is mentioned here")
    m = Matter.create(tmp_path / "m")
    ingest_paths(m, [tmp_path / "f"], None, "c")
    m = Matter(m.root)

    def names(where):
        s = Searcher(Matter(m.root), Options(where=where))
        _, items = s.finish(*s.all())
        return sorted(m.objects.get_object(m.state.metadata[i["loc"][4:]])["name"] for i in items)

    assert names("signature:sqlite") == ["x.db"]                      # part of "SQLite database", any case
    assert names("signature:gzip OR signature:sqlite") == ["x.db", "y.gz"]
    assert names("kind:other -signature:gzip") == ["x.db"]
    assert names("signature:pdf") == []
    # a fields index from before signatures were in it is rebuilt, not refused
    idx = open_index(m, "fields")
    man = idx.manifest()
    man["settings"] = {"schema": 2}
    idx.manifest_path.write_text(_json.dumps(man))
    assert names("signature:sqlite") == ["x.db"]
    assert open_index(m, "fields").manifest()["settings"] == {"schema": 3}


# -- bytes ----------------------------------------------------------------------------------

def test_byte_locators(m):
    doc = doc_named(m, "ledger.txl")
    data = ledger()
    loc = raw.byte_locator(doc, data, 14, 24)
    assert loc["quote"]["hex"] == b"Harrowgate".hex()
    s = locators.format(loc)
    assert s.endswith("#bytes=14,24") and locators.parse(s)["select"] == {"bytes": [14, 24]}
    r = locators.resolve(m, loc)
    assert r["status"] == "exact" and r.data == b"Harrowgate" and r["hex"] == b"Harrowgate".hex()
    bad = {**loc, "quote": raw.fingerprint(b"Bellwether")}
    assert locators.resolve(m, bad)["status"] == "broken"
    assert locators.resolve(m, {"doc": doc, "select": {"bytes": [0, 10 ** 6]}})["status"] == "broken"


def test_hexdump_and_strings():
    rows = raw.hexdump(b"\x00ABC", base=6)
    assert rows[0]["offset"] == 0 and rows[0]["hex"].startswith("                  00 41")
    assert rows[0]["ascii"] == "      .ABC      "
    data = b"\x01\x02" + "Zoë Müller".encode() + b"\x00\x00" + "Windows".encode("utf-16-le") + b"\xff"
    found = raw.strings(data, min_len=5, base=100)
    assert (102, 102 + len("Zoë Müller".encode()), "Zoë Müller", "utf-8") in found
    assert any(t == "Windows" and enc == "utf-16le" for _, _, t, enc in found)


def test_info(m):
    i = raw.info(m, doc_named(m, "ledger.txl"))
    assert i["size"] == len(ledger()) and i["layers"] == [] and not i["utf8"]
    assert any("no text layer" in n for n in i["notes"])
    assert raw.info(m, doc_named(m, "export.json"))["format"] == "json"


# -- pointers and schemas ------------------------------------------------------------------

def test_pointer_ranges():
    text = '{"a": {"b~/c": [10, {"d": "x\\"y"}]}, "e": null}'
    def at(p):
        s, e = structured.pointer_range(text, p)
        return text[s:e]
    assert at("/a/b~0~1c/0") == "10"
    assert at("/a/b~0~1c/1/d") == 'x\\"y'              # the string as written, without its quote marks
    assert at("/e") == "null"
    assert at("") == text
    for bad in ("/a/b~0~1c/2", "/z", "/e/f", "a", "/a/b~0~1c/x"):
        with pytest.raises(structured.PointerError):
            structured.pointer_range(text, bad)
    with pytest.raises(structured.PointerError, match="more than once"):
        structured.pointer_range('{"a": 1, "a": 2}', "/a")
    lines = '{"a": 1}\n\n  {"a": [true]}  \n'
    s, e = structured.pointer_range(lines, "/1/a/0", lines=True)
    assert lines[s:e] == "true"


def test_pointer_locators(m):
    doc = doc_named(m, "export.json")
    r = locators.resolve(m, f"doc:{doc}#ptr=/items/0/note")
    assert r["status"] == "exact" and r["text"] == "Caf\\u00e9 pallet"    # escapes and all: the file's own text
    assert "char" in r["loc"]["select"] and "ptr" not in r["loc"]["select"]
    s = locators.format({"doc": doc, "select": {"ptr": "/a b/c~1d"}})
    assert locators.parse(s)["select"]["ptr"] == "/a b/c~1d"
    ev = doc_named(m, "events.jsonl")
    assert locators.resolve(m, f"doc:{ev}#ptr=/1/msg")["text"] == "signed"
    assert locators.resolve(m, f"doc:{doc_named(m, 'dupkeys.json')}#ptr=/a")["status"] == "broken"
    assert locators.resolve(m, f"doc:{doc_named(m, 'feed.xml')}#ptr=/a")["status"] == "broken"


def test_schema():
    data = {"items": [{"sku": "A", "qty": 1}, {"sku": "B"}]}
    rows = {r["path"]: r for r in structured.schema(data)["paths"]}
    assert rows["/items"]["length"] == 2
    assert rows["/items/*/qty"]["in_objects"] == "1 of 2"
    assert rows["/items/*/sku"]["examples"][1] == {"value": "B", "ptr": "/items/1/sku"}
    recs = structured.schema([{"a": 1}, {"a": "x"}], records_=True)
    assert {r["path"]: r["types"] for r in recs["paths"]}["/*/a"] == {"number": 1, "string": 1}
    x = structured.xml_schema(b'<r><i id="1"><s>A</s></i><i id="2"/></r>')
    assert {p["path"]: p["count"] for p in x["paths"]} == {"/r": 1, "/r/i": 2, "/r/i/@id": 2, "/r/i/s": 1}


# -- interpretations ----------------------------------------------------------------------

def run_ledger(m, name="ledger", **kw):
    return interp.run(m, [doc_named(m, "ledger.txl")], name, [sys.executable, str(m.workspace / "decode.py")],
                      "Ledger records", about="u32 time, u16 length, name, f64 amount", **kw)


def test_run_and_check(m):
    rec = run_ledger(m)
    assert rec["output"]["type"] == "json"
    assert rec["decoder"]["files"][0]["path"] == "decoder/decode.py"
    v = interp.value_of(m, rec, "/records/1")
    assert v["party"] == "Bellwether" and v["_bytes"] == [32, 56]
    assert interp.check(m, "ledger") == {"name": "ledger", "ok": True, "problems": [], "reproducible": True}
    # the decoder ran from its copy: editing the original changes nothing recorded
    (m.workspace / "decode.py").write_text(DECODER.replace('"party"', '"who"'))
    assert interp.check(m, "ledger")["ok"]
    with pytest.raises(interp.InterpError, match="exists"):
        run_ledger(m)
    (interp.folder(m, "ledger") / "decoder" / "decode.py").write_text("tampered")
    assert interp.check(m, "ledger", rerun=False)["problems"] == ["decoder/decode.py has changed since it ran"]
    with pytest.raises(interp.InterpError, match="failed"):
        interp.run(m, [doc_named(m, "ledger.txl")], "bad", [sys.executable, "-c", "raise SystemExit(3)"], "x")
    assert not interp.folder(m, "bad").exists()


def test_rendering(m):
    rec = run_ledger(m)
    html_ = interp.figure_html(m, rec, "/records", "table")
    assert "Interpretation by Claude, not the document’s own content" in html_
    assert "<th>party</th>" in html_ and "_bytes" not in html_
    assert f'loc:doc:{rec["inputs"][0]["doc"]}#bytes=8,32' in html_
    assert interp.shown_as(rec, interp.value_of(m, rec, "/records"), "table", block=False)[1]
    assert interp.default_as(rec, interp.value_of(m, rec), True) == "tree"
    assert interp.as_text(m, rec, interp.value_of(m, rec, "/records/0/party"), "value") == "Harrowgate"


def write_page(path, body, refs):
    path.write_text(body + "\n" + literate.format_refs(refs), encoding="utf-8")
    return path


def test_page_with_bytes_pointers_and_interpretations(m, tmp_path):
    rec = run_ledger(m)
    led, exp = doc_named(m, "ledger.txl"), doc_named(m, "export.json")
    p = write_page(tmp_path / "p.md", "---\ntitle: T\n---\nMagic :q{#magic}. Second party :interp{#p2}.\n\n"
                   "::interp{#recs}\n\n::q{#rec}\n\nAccount :q{#acct}.\n", [])
    pages.add_ref(m, p, "magic", [f"doc:{led}#bytes=0,4"])
    pages.add_ref(m, p, "rec", [f"doc:{led}#bytes=8,32"])
    pages.add_ref(m, p, "acct", [f"doc:{exp}#ptr=/account"])
    pages.add_interp_ref(m, p, "recs", "ledger", "/records")
    pages.add_interp_ref(m, p, "p2", "ledger", "/records/1/party")
    page = pages.Page(m, p)
    assert literate.problems(page.check()) == []
    out, _ = pages.render(page)
    assert "54 58 4c 47" in out and "ledger at bytes 0–3" in out
    assert "|rrowgate" in out                                        # the hex dump's text column
    assert "PO-BMS-26-0112</span>" in out
    assert "Bellwether</span> <span class=\"adhoc-mark\">(Claude’s reading of ledger)" in out
    assert "<th>party</th>" in out
    assert [d for d, _ in pages.cited_docs(page)] == [led, exp]

    prov = pages.export(page, "html", tmp_path / "p.html")
    data = json.loads((tmp_path / "p.html.provenance.json").read_text())
    assert [q["id"] for q in data["quotations"]] == ["magic", "rec", "acct"]
    assert [i["id"] for i in data["interpretations"]] == ["p2", "recs"]
    assert data["interpretations"][0]["unofficial"] and data["interpretations"][0]["rendered"] == "Bellwether"
    assert prov["reanchored"] == []
    pages.export(page, "docx", tmp_path / "p.docx")
    import docx
    d = docx.Document(str(tmp_path / "p.docx"))
    assert "Interpretation by Claude" in "\n".join(x.text for x in d.paragraphs)
    assert [c.text for c in d.tables[0].rows[0].cells] == ["time", "party", "amount"]


def test_interpretations_are_never_quotations(m, tmp_path):
    run_ledger(m)
    led = doc_named(m, "ledger.txl")
    p = write_page(tmp_path / "p.md", "A :q{#recs}. B :interp{#magic}.\n", [])
    pages.add_interp_ref(m, p, "recs", "ledger", "/records/0/party")
    pages.add_ref(m, p, "magic", [f"doc:{led}#bytes=0,4"])
    res = {r["id"]: r for r in pages.Page(m, p).check() if "id" in r and "lint" not in r}
    assert res["recs"]["status"] == "broken" and "can't be quoted" in res["recs"]["error"]
    assert res["magic"]["status"] == "broken" and "not an interpretation" in res["magic"]["error"]
    with pytest.raises(pages.PageError):
        pages.export(pages.Page(m, p), "html", tmp_path / "p.html")


def test_changed_interpretation_breaks_references(m, tmp_path):
    run_ledger(m)
    p = write_page(tmp_path / "p.md", "::interp{#recs}\n", [])
    pages.add_interp_ref(m, p, "recs", "ledger", "/records")
    (m.workspace / "decode.py").write_text(DECODER.replace('"party"', '"who"'))
    run_ledger(m, replace=True)
    r = [x for x in pages.Page(m, p).check() if x.get("id") == "recs"][0]
    assert r["status"] == "broken" and "made again" in r["error"]


def test_png_and_csv_outputs(m, tmp_path):
    png = tmp_path / "png.py"
    png.write_text("import sys, pymupdf\npix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 4, 4), 0)\n"
                   "pix.clear_with(200)\nsys.stdout.buffer.write(pix.tobytes('png'))\n")
    rec = interp.run(m, [doc_named(m, "ledger.txl")], "pic", [sys.executable, str(png)], "A picture")
    assert rec["output"]["type"] == "png"
    assert interp.shown_as(rec, None, None, True) == ("image", None)
    csvp = tmp_path / "csv.py"
    csvp.write_text("print('a,b'); print('1,2')\n")
    rec = interp.run(m, [doc_named(m, "ledger.txl")], "tab", [sys.executable, str(csvp)], "A table", type_="csv")
    assert interp.table_rows(rec, interp.value_of(m, rec))[:2] == (["a", "b"], [["1", "2"]])
    p = write_page(tmp_path / "p.md", "::interp{#pic}\n\n::interp{#tab}\n", [])
    pages.add_interp_ref(m, p, "pic", "pic")
    pages.add_interp_ref(m, p, "tab", "tab")
    pages.export(pages.Page(m, p), "docx", tmp_path / "p.docx")


def test_search_interpretations(m, tmp_path):
    run_ledger(m)
    led = doc_named(m, "ledger.txl")
    header, items = interp.search(m, "bellwether")
    assert header["of"] == "interpretations" and header["exhaustive"]
    [it] = items
    assert it["unofficial"] and it["interp"] == "ledger"
    [hit] = it["hits"]
    assert hit["ptr"] == "/records/1/party" and hit["record"] == "/records/1"
    assert hit["evidence"] == f"doc:{led}#bytes=32,56"                # quote this, not the hit
    assert interp.search(m, "harrowgate /p 1250")[1]                  # /p: the same record
    assert not interp.search(m, "harrowgate /p 42")[1]
    assert not interp.search(m, "party")[1]                           # keys aren't searched
    with pytest.raises(interp.InterpError, match="field"):
        interp.search(m, "kind:other")
    csvp = tmp_path / "csv.py"
    csvp.write_text("print('who,what'); print('dk,signed for pallet 3')\n")
    interp.run(m, [led], "tab", [sys.executable, str(csvp)], "A table", type_="csv")
    header, items = interp.search(m, "pallet", only=["tab"])
    assert header["coverage"]["searched"] == 1
    assert items[0]["hits"][0]["row"] == 1 and items[0]["hits"][0]["column"] == "what"


def test_interp_search_as_a_list(m, tmp_path):
    import base64
    from adhocdocs import lists, view
    run_ledger(m)
    led = doc_named(m, "ledger.txl")
    csvp = tmp_path / "csv.py"
    csvp.write_text("print('who,what'); print('dk,signed for pallet 3')\n")
    interp.run(m, [led], "tab", [sys.executable, str(csvp)], "A table", type_="csv")
    header, items = interp.search(m, "bellwether OR pallet")
    head, rows = interp.search_list(header, items)
    assert head["unofficial"] and head["made_by"] == "interp search"
    by = {r["fields"]["reading"]: r for r in rows}
    assert by["ledger"]["loc"] == f"doc:{led}#bytes=32,56"            # the evidence, not the hit
    assert by["ledger"]["fields"]["where"] == "/records/1/party"
    assert by["tab"]["loc"] == f"doc:{led}" and by["tab"]["fields"]["evidence"].startswith("none")
    assert by["tab"]["fields"]["where"] == "row 1, what"
    assert {r["group"] for r in rows} == {"Ledger records (ledger)", "A table (tab)"}
    p = tmp_path / "r.list.jsonl"
    lists.write(p, head, rows)
    files, seq = view.list_html(m, p)
    page = base64.b64decode(files["index.html"]).decode()
    assert "Interpretation by Claude" in page and "Bellwether" in page
    assert seq[0].startswith("doc:")
